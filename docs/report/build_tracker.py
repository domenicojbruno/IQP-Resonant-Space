from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

PAGES = {
    "Title page": "cover", "Abstract": "i", "Acknowledgements": "ii", "Executive Summary": "iii",
    "Table of Contents": "1-2", "1 Introduction": "3", "2.1": "4", "2.2": "4", "2.3": "4-5", "2.4": "5",
    "2.5": "5", "2.5.1": "5-6", "2.5.2": "6", "3.1": "7", "3.2": "7", "3.2.1": "7", "3.2.2": "7-8",
    "3.2.3": "8", "3.3": "8-9", "3.3.1": "9-11", "3.3.2": "11-12", "3.4": "13", "3.4.1": "13",
    "3.4.2": "14-16", "3.4.3": "16-18", "4 Findings": "19", "4.1": "19-20", "4.2": "20-21",
    "5 Recommendations": "21", "5.1": "21", "5.1.1": "21-22", "5.1.2": "22", "5.1.3": "22",
    "5.1.4": "23", "5.2": "23", "5.2.1": "23-24", "5.2.2": "24", "5.3 Long-Term": "24", "5.3.1": "25",
    "5.3.2": "25", "5.3 Deliverables": "25-26", "5.4 Conclusions": "26", "Bibliography": "27-28",
    "Appendices": "29-31", "Whole document": "all",
}

C, S, R, G = "Content update", "Structure", "Citation", "Grammar / wording"

changes = [
 # ---------------- Content updates ----------------
 ("Title page", C, "Lists one author (Corwin M. Loomis-Price) and the date May 4, 2026, but the text says \"we\" and \"the team.\"", "Update the author list, advisor line, and submission date for the current term. Confirm with your advisor whether this is one continuing report or a new report that cites the earlier one."),
 ("Abstract", C, "Describes only research and a proposed display. No mention of the control system.", "Add one or two sentences on what was built: the actuated wall, the ESP32 motor driver boards, time-of-flight sensing, and the wireless master control software."),
 ("Executive Summary", C, "3. Methodology says the work was research, brainstorming and fabrication only.", "Add the controls work: custom motor driver board, firmware, wireless link, master software, sensor testing."),
 ("Executive Summary", C, "4. Findings: \"the complete immersive installation was not finalized\" and treats the wall as a static prototype.", "Restate what now works and what does not. Use the same status wording as the new Section 4.3."),
 ("Executive Summary", C, "5. Recommendations lists \"Sensor testing and responsive system integration\" as future work.", "Remove or narrow items that are now done; replace with the remaining work (camera tracking, sensor port to ESP32, board revision, lighting, sound, participant study)."),
 ("3.2.3", C, "Methods paragraph mentions \"motion sensors\" in general terms.", "Name the actual sensing approach (VL53L0X time-of-flight threshold sensors, planned overhead cameras) and point to the new Section 3.4.9."),
 ("3.4", C, "Intro says this iteration \"primarily focused on intellectual development and construction of the foundational arches wall structures. Additional features remain areas for continued research.\"", "Rewrite to introduce both halves of the build: mechanical (3.4.1-3.4.3) and electrical/controls (new 3.4.4-3.4.9)."),
 ("3.4.2", C, "Describes the linear actuators but not how they are powered or driven.", "Add one sentence that each actuator is driven by a BTS7960 H-bridge and refer to new Section 3.4.5. Add actuator model, stroke and voltage if known."),
 ("3.4.3", C, "Describes \"arches\" and scissor wall units. Your control spec describes 13 hexagonal actuator positions in three rows (5 / 3 / 5).", "Reconcile the two descriptions. State the actual number of actuators and their layout, and add a layout figure."),
 ("4 Findings", C, "Opening paragraph says the iteration \"focused on conceptual development and construction of the foundational archway wall prototype.\"", "Update to include control system development and bench testing. Keep the \"preliminary\" caveat for wellness outcomes."),
 ("4.1", C, "\"Even without the full integration of responsive systems...\"", "Update to reflect the level of integration actually reached at submission."),
 ("4.2", C, "Six-item list of future focus areas includes \"Implementing motion-responsive interaction.\"", "Move completed items into findings; keep only work that is still open."),
 ("5.1.1", C, "\"Code Base Sensor Testing\" is written as future work for the next team.", "Replace with the new Section 5.1.1 draft (see Sections to Add). Most of this is now current work."),
 ("5.1.2", C, "\"Sensor Integration and Build Finalization\" is written as future work.", "Rewrite as remaining integration steps: install new board, verify pinout, calibrate each actuator, mount sensors."),
 ("5.3 Deliverables", C, "Lists \"Preliminary coding frameworks for responsive interaction systems and their connected electrical systems\" with no detail anywhere in the report.", "Replace with specific deliverables: motor board firmware, gateway firmware, master control application, motor driver board schematic and PCB layout, sensor test sketch, code repository."),
 ("5.4 Conclusions", C, "Two sentences only.", "Expand to a full paragraph covering both the wellness argument and the working technical platform."),
 ("Appendices", C, "Only image appendices A and B.", "Add technical appendices (pinout, command protocol, code repository). See Sections to Add."),
 # ---------------- Structure ----------------
 ("5.3 Deliverables", S, "Section number 5.3 is used twice (Long-Term Recommendations and Deliverables).", "Renumber Deliverables to 5.4 and Conclusions to 5.5."),
 ("Executive Summary", S, "All three Executive Summary pages are numbered III.", "Number them iii, iv, v."),
 ("Table of Contents", S, "Table of Contents pages are numbered 1-2 and the Introduction starts on 3.", "Keep roman numerals through the front matter and start the Introduction on page 1."),
 ("Table of Contents", S, "Entries will be out of date after renumbering and new sections.", "Regenerate the table of contents last."),
 ("Executive Summary", S, "\"Conclusion\" is unnumbered while the other parts are numbered 1-5.", "Number it 6 or remove the numbering from all parts."),
 ("3.1", S, "Headings 3.1-3.4 end in colons (\"Purpose:\", \"Objectives list:\", \"Design Process:\", \"Construction Process:\"); no other headings do.", "Remove the colons. Change \"Objectives list\" to \"Objectives.\""),
 ("3.2.1", S, "Objective 1 has an intro paragraph; Objectives 2 and 3 do not.", "Add a one-sentence intro to Objectives 2 and 3, or remove it from Objective 1."),
 ("Whole document", S, "The wall is called \"arches wall,\" \"arch wall,\" \"archway wall,\" and \"wall prototype.\"", "Pick one term and use it everywhere."),
 ("Whole document", S, "\"Well being,\" \"wellbeing,\" and \"well-being\" are all used.", "Use \"well-being\" throughout."),
 ("Whole document", S, "\"IQP,\" \"Qualifying Projects,\" and \"continuation projects\" are used interchangeably.", "Define IQP once, then use it consistently."),
 ("Whole document", S, "Mixed first person (\"our research,\" \"our community\") and third person (\"the team,\" \"the project\").", "Choose one voice."),
 # ---------------- Citations ----------------
 ("1 Introduction", R, "\"(WPI Well Being Task Force, 2022)\" is cited but missing from the Bibliography.", "Add the full reference."),
 ("3.2.1", R, "\"(Watson, 1988)\" is cited for PANAS, but the listed reference is a different Watson paper (\"Intraindividual and interindividual analyses...\").", "Verify. PANAS is normally cited as Watson, Clark & Tellegen (1988), Journal of Personality and Social Psychology, 54(6), 1063-1070."),
 ("1 Introduction", R, "\"(Chen et al., 2021)\" - the reference has two authors.", "\"(Chen & Zeng, 2021)\""),
 ("2.1", R, "\"(Liu et al, 2024)\" - the reference has two authors.", "\"(Liu & Lin, 2024)\""),
 ("2.5.2", R, "\"(Pinter, 2025)\" - the reference has three authors.", "\"(Pinter et al., 2025)\""),
 ("2.5.1", R, "\"Positive and Negative Affect Surveys (PANAS)\"; also \"Survey\" in 5.1.3.", "\"Positive and Negative Affect Schedule (PANAS)\" in both places."),
 ("5.3.1", R, "\"Worcester Museum of Art\"", "\"Worcester Art Museum\""),
 ("Bibliography", R, "Not in alphabetical order (Hadavi is listed before Godley).", "Move Godley above Hadavi."),
 ("Bibliography", R, "Pinter entry uses a different format from the others (no commas after surnames, \"doi:\" prefix, year repeated).", "Reformat to match the other APA entries."),
 ("Bibliography", R, "No references for the hardware.", "Add datasheets and documentation cited in the new technical sections (see Sections to Add, Bibliography additions)."),
 ("2.3", R, "Second and third paragraphs make research claims with no citation.", "Add a source for the participatory art claims."),
 # ---------------- Grammar: Abstract & Exec Summary ----------------
 ("Abstract", G, "financial security fatigue", "financial insecurity"),
 ("Abstract", G, "within the technically focused university environment", "within a technically focused university environment"),
 ("Executive Summary", G, "personal reflection and interpersonal interaction", "personal reflection, and interpersonal interaction"),
 ("Executive Summary", G, "Research list: commas after the first three items, nothing after the last", "Remove the trailing commas"),
 ("Executive Summary", G, "participant proximity intended to serve", "participant proximity, intended to serve"),
 ("Executive Summary", G, "foundational arches wall prototype", "foundational arch wall prototype"),
 # ---------------- Grammar: Introduction ----------------
 ("1 Introduction", G, "(Barbayannis et al. 2022)", "(Barbayannis et al., 2022)"),
 ("1 Introduction", G, "(Li .et al, 2025)", "(Li et al., 2025)"),
 ("1 Introduction", G, "academic years, where students surveyed reported", "academic years, when surveyed students reported"),
 ("1 Introduction", G, "wellness/ mental health days", "wellness days"),
 ("1 Introduction", G, "peak in 2021-22 many students", "peak in 2021-22, many students"),
 ("1 Introduction", G, "wielded as effective weapons in the fight to combat negative mental health", "used effectively to address poor mental health"),
 ("1 Introduction", G, "Furthermore the engineered artistic expression aligns ... institution inspiring", "Furthermore, engineered artistic expression aligns ... institution, inspiring"),
 ("1 Introduction", G, "student well being", "student well-being"),
 ("1 Introduction", G, "The project focus was on gaps in the student wellness infrastructure evaluating", "The project focused on gaps in the student wellness infrastructure, evaluating"),
 ("1 Introduction", G, "mechanical creation the project aimed", "mechanical creation, the project aimed"),
 ("1 Introduction", G, "heralding the student mental wellness initiative", "supporting the student mental wellness initiative"),
 # ---------------- Grammar: Background ----------------
 ("2.1", G, "(Li et al, 2025)", "(Li et al., 2025)"),
 ("2.1", G, "reduced intrapersonal connectedness", "reduced interpersonal connectedness"),
 ("2.1", G, "(Zarowski et al., 2024) leading to", "(Zarowski et al., 2024), leading to"),
 ("2.2", G, "to support students including:", "to support students, including:"),
 ("2.2", G, "List items: \"Implemented community-centered spaces\" / \"Promote peer support initiatives\"", "\"Community-centered spaces\" / \"Peer support initiatives\" (match the other noun items)"),
 ("2.2", G, "necessary important services", "necessary services"),
 ("2.2", G, "improved the situation however issues continue", "improved the situation; however, issues continue"),
 ("2.3", G, "self expression", "self-expression"),
 ("2.3", G, "promote ... connection within a shared sensory experience rather than simply observing passively", "... within a shared sensory experience, rather than passive observation"),
 ("2.3", G, "social confidence and overall wellness", "social confidence, and overall wellness"),
 ("2.3", G, "could be a potentially effective method", "may be an effective method"),
 ("2.4", G, "between engineering artistic expression, technology", "between engineering, artistic expression, technology"),
 ("2.4", G, "In fact it's quite the opposite, when implemented correctly engineering can", "In fact, the opposite is true: when implemented correctly, engineering can"),
 ("2.4", G, "not only as a functional tool but for major emotional and social enhancement", "not only as functional tools but also for emotional and social enhancement"),
 ("2.4", G, "lighting systems and mechanical interaction an immersive", "lighting systems, and mechanical interaction, an immersive"),
 ("2.4", G, "the participant; where the whole", "the participant, where the whole"),
 ("2.4", G, "The interference patterns installation", "The Interference Patterns installation"),
 ("2.4", G, "react between fragmented and cohesive states", "shift between fragmented and cohesive states"),
 ("2.4", G, "In doing so the installation", "In doing so, the installation"),
 ("2.4", G, "very aptly defines", "illustrates"),
 ("2.5", G, "Precedental Art Research", "Precedent Art Research"),
 ("2.5.1", G, "The aforementioned project studied", "The study examined"),
 ("2.5.1", G, "the exhibit where the results indicated an increase in positive affect and decreased negative affect", "the exhibit; the results indicated an increase in positive affect and a decrease in negative affect"),
 ("2.5.1", G, "a fantastic initial framework", "a strong initial framework"),
 ("2.5.1", G, "one's mental response", "participants' mental responses"),
 ("2.5.1", G, "the VisualEars project focus was primarily on", "the VisualEars project focused primarily on"),
 ("2.5.1", G, "expand upon their finding", "expand upon its findings"),
 ("2.5.2", G, "was a physical interactive installation evolving into a powerful broad digital support platform", "was a physical interactive installation that evolved into a broad digital support platform"),
 ("2.5.2", G, "not only for engagement aesthetically but for self expression", "not only for aesthetic engagement but also for self-expression"),
 ("2.5.2", G, "further expands their concept towards", "extends this concept to"),
 ("2.5.2", G, "emotional well being", "emotional well-being"),
 # ---------------- Grammar: Methodology ----------------
 ("3.2", G, "challenges in the student wellness at WPI. Exploring potential solutions of ... (sentence fragment)", "challenges in student wellness at WPI and exploring potential solutions involving ..."),
 ("3.2.1", G, "... within the WPI community. By examining student perceptions of the current wellness ventures. (fragment)", "... within the WPI community by examining student perceptions of current wellness initiatives."),
 ("3.2.1", G, "The study developed by (Watson, 1988)", "The scale developed by Watson (1988)"),
 ("3.2.1", G, "pre and post visiting the installation", "before and after visiting the installation"),
 ("3.2.2", G, "How can interactive installation promote", "How can interactive installations promote"),
 ("3.2.2", G, "feasibility on campus life at WPI", "feasibility on the WPI campus"),
 ("3.2.2", G, "helped better the understanding of", "improved the team's understanding of"),
 ("3.2.3", G, "... participation and reflection. (question ends with a period)", "... participation and reflection?"),
 ("3.2.3", G, "connected to and with the participant", "connected to the participant"),
 ("3.3", G, "the immersive wellness focused installation", "an immersive, wellness-focused installation"),
 ("3.3", G, "the campus communities interests", "the campus community's interests"),
 ("3.3", G, "Rather than function as a traditional one size linear engineering design process", "Rather than functioning as a traditional, linear engineering design process"),
 ("3.3", G, "space for future works / in the Qualifying Projects space", "room for future work / in future Qualifying Projects"),
 ("3.3.1", G, "idea proliferation; whilst providing", "idea generation while providing"),
 ("3.3.1", G, "informal low-pressure environment", "informal, low-pressure environment"),
 ("3.3.1", G, "Throughout the brainstorming session participants", "Throughout the brainstorming sessions, participants"),
 ("3.3.1", G, "Motionally expressive artistic spaces", "Emotionally expressive artistic spaces"),
 ("3.3.1", G, "ambient audioscapes", "ambient soundscapes"),
 ("3.3.1", G, "emphasized the immersion ... while the students from the engineering backgrounds", "emphasized immersion ... while students from engineering backgrounds"),
 ("3.3.1", G, "implementation.The (missing space)", "implementation. The"),
 ("3.3.1", G, "Figure 2 :", "Figure 2:"),
 ("3.3.1", G, "fabrication process becoming central guides through the entire project", "fabrication process, becoming central guides throughout the project"),
 ("3.3.2", G, "The weekly meeting sessions were conducted", "Weekly meetings were held"),
 ("3.3.2", G, "a variety of disciplines frequent communication", "a variety of disciplines, frequent communication"),
 ("3.3.2", G, "practical engineering concentrations", "practical engineering work"),
 ("3.3.2", G, "check points", "checkpoints"),
 ("3.4", G, "foundational arches wall structures", "foundational arch wall structure"),
 ("3.4.1", G, "The parts were printed and intended to allow", "The parts were intended to allow"),
 ("3.4.1", G, "An additional printed housing sheath was printed", "An additional housing sheath was printed"),
 ("3.4.1", G, "a flush fit against the locking each other in place (garbled)", "a flush fit, locking the two in place"),
 ("3.4.2", G, "Following solidworks fabrication and printing the linear actuators", "Following SolidWorks design and printing, the linear actuators"),
 ("3.4.2", G, "... support structures. Where this phase focused on (fragment)", "... support structures. This phase focused on"),
 ("3.4.2", G, "Figure 11: Scissor wall unit assembly integration of the mechanical linkage system", "Figure 11: Scissor wall unit assembly, showing integration of the mechanical linkage system"),
 ("3.4.3", G, "3.4.3Wall Creation (missing space)", "3.4.3 Wall Creation"),
 # ---------------- Grammar: Findings ----------------
 ("4 Findings", G, "timeline of the project the current iteration", "timeline of the project, the current iteration"),
 ("4 Findings", G, "Analysis of ... collaborative engineering and the construction ... was conducted.", "Analyses of ... collaborative engineering, and the construction ... were conducted."),
 ("4 Findings", G, "The support from the project shines through by encouraging", "The project provides this support by encouraging"),
 ("4.1", G, "will need numerous more projects similar to this one; the successful", "will require many more projects like this one, the successful"),
 ("4.1", G, "can be involved and evolved on the campus", "can be combined and developed on campus"),
 ("4.1", G, "Additionally the complete wall", "Additionally, the completed wall"),
 ("4.1", G, "iteration of the project it helped", "iteration of the project, it helped"),
 ("4.1", G, "rather than existing away from it", "rather than existing apart from it"),
 ("4.2", G, "the Inference Patterns: A Resonant Space", "Interference Patterns: A Resonant Space"),
 ("4.2", G, "participatory design session indicated future viability", "participatory design sessions indicated the future viability"),
 ("4.2", G, "engaged students in a similar fashion consistent with", "engaged students in a manner consistent with"),
 ("4.2", G, "the project hypothesis where art through an engineering environment", "the project hypothesis that art in an engineering environment"),
 ("4.2", G, "technical effects structures for participant interaction", "technical systems for participant interaction"),
 ("4.2", G, "significant background or coordination", "significant coordination"),
 ("4.2", G, "longer term fabrication timelines", "longer fabrication timelines"),
 ("4.2", G, "whether repeated engagement added to measurable improvements", "whether repeated engagement leads to measurable improvements"),
 ("4.2", G, "in the scope of this project conceptual development", "within the scope of this project, the conceptual development"),
 ("4.2", G, "in the promotion of student wellness and strengthening campus community", "for promoting student wellness and strengthening the campus community"),
 # ---------------- Grammar: Recommendations ----------------
 ("5 Recommendations", G, "recommendations from preliminary research ... conducted throughout the WPI administration to improve (garbled)", "recommendations, drawn from the preliminary research ... conducted throughout the project, for the WPI administration to improve"),
 ("5 Recommendations", G, "the short and long term recommendations", "short- and long-term recommendations"),
 ("5.1", G, "in the long-term", "in the long term"),
 ("5.1.1", G, "The following IQP continuation team", "The next IQP continuation team"),
 ("5.1.1", G, "the artworks foundational sensor", "the artwork's foundational sensor"),
 ("5.1.1", G, "This variation in motion sensitivity", "Variation in motion sensitivity"),
 ("5.1.2", G, "focus on the responsive systems integration directly into the walls prototype; and finalize", "focus on integrating the responsive systems directly into the wall prototype and finalizing"),
 ("5.1.2", G, "from preliminary into the full functional experience", "from a preliminary prototype into the fully functional experience"),
 ("5.1.3", G, "to denote and store behavioral responses", "to record behavioral responses"),
 ("5.1.3", G, "A created ... (PANAS) as outlined in (Hadavi et al., 2022) should", "A ... (PANAS), as outlined in Hadavi et al. (2022), should"),
 ("5.1.3", G, "for quantifying research goals", "to quantify research outcomes"),
 ("5.1.4", G, "Where each sensory element should be (fragment)", "Each sensory element should be"),
 ("5.2", G, "Short Term Recommendations / short term recommendations", "Short-Term Recommendations / short-term recommendations"),
 ("5.2.1", G, "highly trafficked", "high-traffic"),
 ("5.2.1", G, "library entrance or on quad space near or on the seal", "library entrance, or the Quad near the seal"),
 ("5.2.1", G, "the exhibits reception", "the exhibit's reception"),
 ("5.2.1", G, "Where potential features could include but not be limited to: (fragment)", "Potential features could include, but are not limited to:"),
 ("5.2.1", G, "to compliment the created archway", "to complement the archway"),
 ("5.2.1", G, "Light responsive systems / Student controlled moving display", "Light-responsive systems / Student-controlled moving display"),
 ("5.2.1", G, "gather the feedback", "gather feedback"),
 ("5.2.2", G, "Special attentions to factors including but are not limited to: (fragment)", "Special attention should be given to factors including, but not limited to:"),
 ("5.3 Long-Term", G, "Long-Term Recommendations For The Wellness Space", "Long-Term Recommendations for the Wellness Space"),
 ("5.3 Long-Term", G, "art-engineering-centered", "art- and engineering-centered"),
 ("5.3.1", G, "pilot implementation WPI should", "pilot implementation, WPI should"),
 ("5.3.1", G, "WPI administration after analysis of student and faculty feedback could", "WPI administration, after analysis of student and faculty feedback, could"),
 ("5.3.1", G, "interactive art for students by students", "interactive art for students, by students"),
 ("5.3.1", G, "Unlike the traditional passive art displays like those at ... this permanent", "Unlike traditional passive art displays such as those at ..., this permanent"),
 ("5.3.1", G, "community collaboration encouraging:", "community collaboration, encouraging:"),
 ("5.3.2", G, "This project denotes the potential relationship", "This project indicates a potential relationship"),
 ("5.3.2", G, "However additional research", "However, additional research"),
 ("5.3.2", G, "the longer term effectiveness of this standard of project", "the long-term effectiveness of this type of project"),
 ("5.3.2", G, "which approach is most feasible in promoting", "which approach is most effective in promoting"),
 ("5.3 Deliverables", G, "project produces several deliverables", "project produced several deliverables"),
 ("5.3 Deliverables", G, "provides a strong feasibility of engineering integration in artistic expression for the improved mental wellness", "demonstrates the feasibility of integrating engineering with artistic expression to improve mental wellness"),
 ("5.3 Deliverables", G, "Participatory brainstorming session materials, (only list item with a comma)", "Remove the comma"),
 ("5.3 Deliverables", G, "fully immersed state.Collectively", "fully immersive state. Collectively"),
 ("5.4 Conclusions", G, "interactive artistic engineering focussed environments", "interactive, engineering-focused artistic environments"),
]

PRI = {C: "High", S: "Medium", R: "Medium", G: "Low"}

# ---------------------------------------------------------------------------
# Sections to add
# ---------------------------------------------------------------------------
adds = [
 ("2.6", "Sensing, Actuation, and Control Technology", "Background, after 2.5 Precedent Art Research",
  "The report never explains the technology the wall depends on. A reader needs this before the new construction sections.",
  "Responsive installations depend on three technical capabilities: sensing the presence of a participant, producing physical motion, and coordinating the two. This project used low-cost components that are common in student engineering work so that future teams can maintain and extend the system.\n\n"
  "Presence was sensed with the VL53L0X time-of-flight (ToF) ranging sensor. A ToF sensor emits pulses of infrared laser light and measures the time the reflected light takes to return, which gives an absolute distance that does not depend on the color or reflectance of the target in the way that simple infrared proximity sensors do. The VL53L0X reports distances up to roughly two meters over an I2C connection, which is sufficient to detect a person crossing a hallway threshold.\n\n"
  "Motion was produced by electric linear actuators. A DC actuator reverses direction when the polarity of its supply is reversed, so each actuator requires an H-bridge, a circuit of four switches that can connect the motor to the supply in either direction. The BTS7960 is a high-current half-bridge device; two of them form a full H-bridge that a microcontroller can command with two pulse-width-modulated (PWM) signals, one for each direction, where the duty cycle sets the speed.\n\n"
  "Control was handled by the ESP32-WROOM-32, a microcontroller module with integrated Wi-Fi and enough PWM-capable pins to drive several actuators at once. The ESP32 also supports ESP-NOW, a connectionless wireless protocol from Espressif that lets boards exchange short packets addressed by hardware (MAC) address without a router. This allowed the wall's controller boards to be commanded without signal wiring between them.",
  "Add citations for the three datasheets and the ESP-NOW documentation. Confirm the sensor range figure against your own test results.",
  "DS_vl53l0x.pdf; BTS7960 Motor Driver.pdf; esp32-wroom-32_datasheet_en.pdf"),

 ("3.2.4", "Objective 4: Develop the Responsive Control System", "Methodology, after Objective 3",
  "The three existing objectives cover research and concept only. The controls work needs its own objective so the Findings can report against it.",
  "This objective focused on giving the fabricated wall the ability to respond to people. The team set out to design the electronics, firmware, and software needed to detect a participant and move individual wall sections in response.\n\n"
  "Objective 4: Research Questions:\n"
  "- How can a participant's presence and position be detected reliably in a public hallway?\n"
  "- How can many un-encoded linear actuators be driven safely and repeatably?\n"
  "- How should motion be distributed across the wall so that the response reads as intentional rather than mechanical?\n\n"
  "Methods:\n"
  "The team selected components, designed a custom motor driver board, wrote firmware for the controller boards, and wrote a master control application. Each part was tested on the bench before integration. The construction of these systems is described in Sections 3.4.4 through 3.4.9.",
  "Adjust the research questions to match what your advisor expects. Update Section 3.2 intro to say there are four objectives.",
  "Project description; claude/resonant-wall-master-spec.md"),

 ("3.4.4", "Control System Architecture", "Construction Process, after 3.4.3 Wall Creation",
  "Gives the reader the whole system before the detail sections. Needs one block-diagram figure.",
  "The control system was organized in three layers. At the top, a master control application running on a PC decides how far each section of the wall should extend. In the middle, a single ESP32 connected to the PC by USB acts as a wireless gateway. At the bottom, three ESP32 motor driver boards mounted on the wall each control one horizontal row of actuators.\n\n"
  "The wall was laid out as three rows of actuated sections in a staggered honeycomb pattern: five in the top row, three in the middle row, and five in the bottom row, for thirteen actuators in total. Each actuator is identified by its row and column (for example, M2,3 is the third actuator in the middle row), and each row is owned by one board. Figure [X] shows the layout and the flow of commands from the master to the actuators.\n\n"
  "This arrangement keeps the boards on the wall simple. A board knows only how to move its own actuators safely and report their state. All decisions about which sections move, and when, are made by the master, so the behavior of the installation can be changed in software without reflashing the boards.",
  "Confirm the 5 / 3 / 5 layout and thirteen-actuator count against the physical wall. Add a block diagram and a layout figure.",
  "claude/resonant-wall-master-spec.md; claude/master-ui-and-gateway.md"),

 ("3.4.5", "Motor Driver Board Design", "After 3.4.4",
  "The custom PCB is a major deliverable and is not mentioned anywhere in the draft.",
  "A custom printed circuit board was designed in KiCad to hold one ESP32-WROOM-32 development module and connections for up to five BTS7960 motor drivers. Each driver receives two PWM signals from the ESP32, one that extends the actuator and one that retracts it.\n\n"
  "Safety at power-up was a design requirement. The enable inputs of all five drivers are tied to a single ESP32 pin that is pulled to ground through a 10 kOhm resistor. The drivers therefore remain disabled whenever the microcontroller is unpowered, resetting, or has not yet started its program, and the actuators cannot move until the firmware deliberately enables them.\n\n"
  "Each actuator has one limit-switch input. The extend and retract limit switches for an actuator share that single input, which halves the number of pins required; the firmware determines which switch was reached from the direction the actuator was travelling. The board also reserves pins for five infrared sensor inputs and one addressable LED data line for future lighting. The schematic is shown in Figure [X] and the full pin assignment is given in Appendix C.\n\n"
  "An earlier prototype board, which carried only the PWM connections and no limit switches, was used for initial bench testing while the revised board was being prepared.",
  "State whether the new board has been fabricated and tested by submission. Add power-supply details (voltage, current, connector, fusing). Add schematic and PCB figures.",
  "MotorDriver Board.pdf; MotorDriver Board.kicad_pcb; firmware/motor-board/config.h"),

 ("3.4.6", "Motor Board Firmware", "After 3.4.5",
  "Documents how the actuators are controlled and how position is known without encoders.",
  "A single firmware program, written in C++ with the Arduino framework for ESP32, runs on all three motor driver boards. Two settings identify each board: its row and the number of actuators it drives.\n\n"
  "The linear actuators have no position sensors, so the firmware estimates position from time. During a calibration routine, each actuator is retracted to its limit switch, then driven to the opposite limit while the firmware measures how long the full stroke takes. That travel time is stored in the board's non-volatile memory, separately for each actuator, because travel time varies from one actuator to another. Afterward, a request to move to a given percentage of extension is converted into a drive time.\n\n"
  "At power-up, every actuator performs a homing move to its retract limit so that it starts from a known position. Because the two limit switches share one input, an actuator that starts with a switch already pressed first moves outward briefly and then approaches the retract limit, which guarantees it ends at the correct end.\n\n"
  "Each actuator is managed as a small state machine with states for homing, idle, moving, calibrating, and fault. Every move is supervised by a timeout set to 1.5 times the measured travel time. If no limit switch is reached within that time, the firmware stops the actuator and reports a fault, which protects the mechanism if a switch fails. PWM is generated at 20 kHz, above the range of human hearing, so the drivers do not produce an audible tone in a space intended to be calming.",
  "Add measured travel times once calibration has been run on the wall. Confirm the PWM frequency and timeout values are the ones used at submission.",
  "firmware/motor-board/README.md; MotorController.cpp; config.h"),

 ("3.4.7", "Wireless Communication", "After 3.4.6",
  "Explains the ESP-NOW link and why it replaced the serial plan.",
  "The original plan was to command the boards over a wired serial connection. This was replaced with ESP-NOW, which requires no signal wiring between the master and the wall and addresses each board directly by its hardware address.\n\n"
  "Commands are sent as four-byte packets containing an identifying byte, the actuator number, the command, and a value. Six commands are defined: drive forward, drive in reverse, stop, home, calibrate, and go to a position given as a percentage of full extension. A reserved actuator number applies a command to every actuator on a board.\n\n"
  "Each board reports back to the master four times per second with a six-byte status packet for every actuator, giving its state, estimated position, and any fault. An emergency stop is implemented by broadcasting a stop command to all boards at once and then repeating it to each board individually. The complete packet formats are listed in Appendix D.\n\n"
  "Each board also provides a text console over its USB port that accepts the same commands, which allowed individual boards to be tested without the rest of the system.",
  "Report measured range and reliability in the installation location if you test it. Confirm the status rate (250 ms) is unchanged.",
  "claude/master-ui-and-gateway.md; firmware/motor-board/CommsEspNow.cpp; firmware/gateway/ResonantWallGateway.ino"),

 ("3.4.8", "Master Control Software and Interaction Logic", "After 3.4.7",
  "This is where the engineering connects to the artistic intent. Currently absent from the report.",
  "The master control application was written in Python. It provides a graphical interface for configuring each board's connection, sending manual commands, monitoring the state of every actuator, and stopping the wall.\n\n"
  "The interaction model treats a participant as a point on the face of the wall. Each wall section calculates its distance from that point, and its target extension is reduced according to a Gaussian (bell-shaped) function of that distance. The section directly in front of the participant retracts fully, neighboring sections retract partway, and distant sections remain extended. The result is a smooth recess that follows the participant along the wall, rather than individual sections switching on and off. Two parameters shape the behavior: a radius that sets how far the response spreads to neighboring sections, and a smoothing rate that sets how quickly sections follow their targets.\n\n"
  "To prevent sections from chattering when a participant stands still near a boundary, commands are issued with hysteresis: a section is commanded to retract only when its target falls below 12 percent extension and to extend again only when its target rises above 85 percent. When the hallway is unoccupied, all sections return to full extension.\n\n"
  "This behavior was designed to express the installation's theme. The wall yields to a person's presence and closes behind them, and the shared response of neighboring sections reflects the project's emphasis on connection rather than isolation.",
  "Add a screenshot of the interface. State which mode (threshold or position streaming) was used on the final wall. Adjust the final paragraph to match the team's own artistic intent.",
  "master/resonant_wall_master.py; claude/resonant-wall-master-spec.md"),

 ("3.4.9", "Occupancy and Position Sensing", "After 3.4.8",
  "Documents the ToF sensor work, which is the core of the project description.",
  "Time-of-flight sensors were chosen to detect when a person enters the space in front of the wall. The design calls for VL53L0X sensors positioned at the hallway thresholds; a distance reading shorter than the width of the hallway indicates that someone has crossed.\n\n"
  "The sensor was first evaluated on an Arduino Uno. It was configured in its long-range mode, which trades some speed and accuracy for greater detection distance, and polled ten times per second. The test confirmed that the sensor was suitable for threshold detection at hallway distances. Because the sensor measures reflected infrared light, its range is expected to be reduced in direct sunlight, which should be considered when the installation site is chosen.\n\n"
  "When several of these sensors share one I2C bus, each must be assigned a unique address at start-up. This is done by holding all sensors in shutdown and enabling them one at a time, assigning each an address before the next is enabled.\n\n"
  "The threshold sensors establish whether the space is occupied. Tracking where a participant is along the wall is planned to use overhead cameras, with the tracked position supplied to the interaction logic described in Section 3.4.8.",
  "Add your measured detection distances. State how many sensors were installed, where, and whether they are read by an ESP32 at submission. State the status of camera tracking honestly (planned, prototyped, or working).",
  "claude/vl53l0x-arduino-uno-test.md; TOFpins.png; DS_vl53l0x.pdf"),

 ("3.5", "Safety Considerations", "Methodology, after 3.4",
  "A moving wall in a public hallway raises obvious safety questions that reviewers will ask. The report says nothing about them.",
  "Because the installation moves in a space shared with the public, safety was treated as a design constraint. Several measures were built into the electronics and firmware. The motor drivers are disabled by default in hardware and are enabled only after the firmware has started. Each actuator stops at its limit switches, and a timeout stops any actuator that fails to reach one. An emergency stop in the master application halts every actuator with a single action, and closing the application sends the same stop command.\n\n"
  "Some risks remain in the current design. The scissor mechanisms create pinch points that are not yet guarded. The boards cannot measure motor current, so an actuator that is blocked by an obstacle will continue to push until its timeout expires. [Describe the force the actuators can apply and any physical guarding, speed limits, or signage used.] These items should be addressed before the installation is left unattended in a public location.",
  "Fill in the bracketed sentence. Check with your advisor whether WPI requires a safety review or IRB approval for participant studies.",
  "firmware/motor-board/README.md; claude/motor-board-peripheral-firmware.md"),

 ("4.3", "Control System Development and Bench Testing", "Findings, after 4.2",
  "Findings currently report only on the mechanical wall and the brainstorming sessions.",
  "The control system was developed and tested in stages. The time-of-flight sensor was confirmed to detect targets at hallway distances when tested on an Arduino Uno. The motor board firmware was exercised on the prototype driver board, where [N] actuators were driven in both directions from the USB console. The master control application was verified against simulated boards that reproduced the firmware's command and status packets, which confirmed address detection, delivery acknowledgement, status decoding, both motion modes, and the emergency stop.\n\n"
  "[Report results on real hardware here: whether the gateway and boards communicated over ESP-NOW, whether homing and calibration ran on the revised board, the measured travel time of each actuator, and whether the wall responded to a person.]\n\n"
  "These results show that a wall of independently actuated sections can be coordinated wirelessly from a single controller using inexpensive components, and that the interaction behavior can be developed and adjusted in software separately from the hardware.",
  "IMPORTANT: replace both bracketed passages with what was actually tested. As of the project notes, the revised board had not been built and the gateway had not been run on hardware. Do not keep the last paragraph unless the hardware results support it.",
  "claude/master-ui-and-gateway.md (Status); claude/vl53l0x-arduino-uno-test.md"),

 ("4.4", "Technical Limitations", "Findings, after 4.3",
  "Stating limits plainly strengthens the report and gives the next team a starting list.",
  "Several limitations of the current system should be noted. Position is estimated from drive time rather than measured, so accuracy depends on calibration and will drift if the load on an actuator or the supply voltage changes; each actuator re-establishes its position only when it reaches a limit switch. The boards have no current sensing, so a stalled or obstructed actuator is detected only by timeout. On the revised board, two limit-switch inputs share pins with the USB serial port, which restricts use of the USB console while those actuators are running. The infrared sensor inputs and LED output on the board are not yet used.\n\n"
  "Most importantly, no participant data were collected. The system's effect on mood, stress, or social connection has not been measured, and the wellness claims in this report remain a hypothesis supported by the literature rather than by this installation.",
  "Remove any limitation that has been fixed by submission.",
  "claude/motor-board-peripheral-firmware.md (Open items)"),

 ("5.1.1 (replacement)", "Complete Hardware Integration and Calibration", "Replaces existing 5.1.1 Code Base Sensor Testing",
  "The existing text describes work that is now largely done.",
  "The next team should begin by completing integration of the control system on the wall. The pin assignments of the revised motor driver board should be verified against the fabricated board before motors are connected. Each board should then be given its identity and the gateway's address, and every actuator should be homed and calibrated so that its travel time is stored. The team should confirm that limit switches stop each actuator at both ends and that the emergency stop halts the whole wall.\n\n"
  "Two hardware changes are recommended for the next board revision: moving the two limit-switch inputs that share the USB serial pins to other pins, and connecting the motor drivers' current-sense outputs to the ESP32 so that stalls and obstructions can be detected directly.",
  "Delete steps that are complete by submission.",
  "firmware/motor-board/README.md (Suggested next steps)"),

 ("5.1.2 (replacement)", "Sensor Integration and Participant Tracking", "Replaces existing 5.1.2",
  "Narrows the old section to what is still open.",
  "With the wall moving reliably, the next step is to connect it to people. The time-of-flight sensor code should be moved from the Arduino Uno to an ESP32, extended to several sensors on one bus, and used to signal the master when the hallway is occupied. The camera-based tracking that supplies a participant's position along the wall should then be developed and mapped onto the wall's coordinate system. The interaction parameters, including the response radius, smoothing rate, and thresholds, should be tuned with real participants rather than chosen in advance.\n\n"
  "The lighting and sound elements proposed during the participatory sessions can then be layered onto the same master application, using the LED output already present on the boards.",
  "Confirm the status of camera tracking before using this text.",
  "claude/vl53l0x-arduino-uno-test.md (Open items); claude/resonant-wall-master-spec.md"),

 ("Appendix C", "Motor Driver Board Pin Assignments", "After Appendix B",
  "Future teams need this to wire and debug the boards.",
  "Table C1. ESP32 pin assignments on the motor driver board (revised board).\n\n"
  "Driver enable (all five drivers): GPIO 2, pulled down through 10 kOhm\n"
  "Actuator 0 extend / retract PWM: GPIO 12 / 13; limit input: GPIO 3\n"
  "Actuator 1 extend / retract PWM: GPIO 26 / 25; limit input: GPIO 1\n"
  "Actuator 2 extend / retract PWM: GPIO 32 / 27; limit input: GPIO 21\n"
  "Actuator 3 extend / retract PWM: GPIO 18 / 19; limit input: GPIO 14\n"
  "Actuator 4 extend / retract PWM: GPIO 22 / 23; limit input: GPIO 15\n"
  "Infrared sensor inputs (unused): GPIO 36, 39, 34, 35, 5\n"
  "LED strip data (unused): GPIO 4\n"
  "Status LED: GPIO 16",
  "This table was transcribed from the schematic PDF and is marked unverified in the project notes. Check every row against the built board before publishing. Format as a real table in the report.",
  "firmware/motor-board/config.h; claude/motor-board-peripheral-firmware.md"),

 ("Appendix D", "Command and Status Protocol", "After Appendix C",
  "Documents the interface so the master and boards can be developed independently.",
  "Command packet (master to board), 4 bytes: [0xA5, actuator index, command, value]. An actuator index of 0xFF addresses every actuator on the board.\n\n"
  "Commands: 0 forward, 1 reverse, 2 stop, 3 home, 4 calibrate, 5 go to position. The value is a PWM duty of 0-255 for forward and reverse, a percentage of 0-100 for go to position, and is ignored otherwise.\n\n"
  "Status packet (board to master), 6 bytes: [0xA6, board number, actuator index, state, position percent, fault], sent every 250 ms for each actuator. A position of 255 means the position is not yet known.\n\n"
  "States: 0 uninitialized, 1 homing (clearing limit), 2 homing (seeking limit), 3 idle, 4 moving forward, 5 moving reverse, 6 calibrating, 7 fault.\n"
  "Faults: 0 none, 1 homing timeout, 2 runaway timeout.",
  "Format as tables. Confirm the values match the firmware at submission.",
  "claude/master-ui-and-gateway.md (Protocol)"),

 ("Appendix E", "Source Code and Design Files", "After Appendix D",
  "Lets future teams find the code. Standard for a continuation project.",
  "All firmware, software, and board design files produced for this project are available at https://github.com/domenicojbruno/IQP-Resonant-Space. The repository contains the motor board firmware, the gateway firmware, the master control application, and the KiCad files for the motor driver board. Instructions for building and flashing the firmware are included in the repository.",
  "Confirm the repository is public (or accessible to WPI) and that the listed contents are actually committed.",
  "Project instructions; claude/github-repo-index.md"),

 ("Front matter", "List of Figures", "After Table of Contents",
  "The report has 16 figures and will gain more. A figure list is standard in WPI project reports.",
  "List of Figures\n\nFigure 1: Participatory brainstorming session ... [page]\nFigure 2: Early collaborative concept ideation sessions ... [page]\n[Continue for every figure. Generate automatically from the captions in your word processor.]",
  "Generate automatically; do not type by hand.",
  "Existing figure captions"),

 ("Front matter", "Authorship", "After Acknowledgements",
  "The text says \"we\" but one author is listed. An authorship statement makes clear who wrote and built what, especially for a continuation.",
  "[Name] wrote Sections [..] and led [mechanical design / fabrication]. [Name] wrote Sections [..] and led [electronics / firmware / software]. All authors reviewed and edited the full report.",
  "Fill in. Check whether your advisor requires this section.",
  "None - needs your input"),

 ("Bibliography", "Bibliography additions", "Merge alphabetically into the existing list",
  "The new technical sections cite hardware that has no references, and one existing citation has no entry.",
  "Espressif Systems. ([year]). ESP32-WROOM-32 datasheet (Version [x]). [URL]\n\n"
  "Espressif Systems. ([year]). ESP-NOW. ESP-IDF Programming Guide. [URL]\n\n"
  "Infineon Technologies. ([year]). BTS 7960 high current PN half bridge NovalithIC (Rev. [x]). [URL]\n\n"
  "STMicroelectronics. ([year]). VL53L0X: Time-of-Flight ranging sensor datasheet (Rev. [x]). [URL]\n\n"
  "Watson, D., Clark, L. A., & Tellegen, A. (1988). Development and validation of brief measures of positive and negative affect: The PANAS scales. Journal of Personality and Social Psychology, 54(6), 1063-1070.\n\n"
  "WPI Mental Health and Well-Being Task Force. (2022). [Report title]. Worcester Polytechnic Institute. [URL]",
  "Take the year and revision from each datasheet's cover page. Verify the Watson reference and the exact title of the WPI task force report before using them.",
  "Datasheet PDFs in the project"),
]

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
FONT = "Arial"
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="1F3A4D")
body = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D5DA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUSES = '"To do,In progress,Done,Not needed"'

def header(ws, cols, widths, row=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

def status_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", bgColor="CFE8D5")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"In progress"'], fill=PatternFill("solid", bgColor="FCE9B8")))

wb = Workbook()

# ---- Sheet: Changes by Section ----
ws = wb.active
ws.title = "Changes by Section"
cols = ["#", "Section", "Page", "Type", "Current text / issue", "Change to make", "Priority", "Status", "Notes"]
header(ws, cols, [5, 20, 8, 17, 58, 58, 10, 13, 28])
for n, (sec, typ, cur, fix) in enumerate(changes, 1):
    row = [n, sec, PAGES[sec], typ, cur, fix, PRI[typ], "To do", ""]
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font, c.alignment, c.border = body, wrap, border
last = ws.max_row
ws.auto_filter.ref = f"A1:I{last}"
dv = DataValidation(type="list", formula1=STATUSES, allow_blank=True)
ws.add_data_validation(dv); dv.add(f"H2:H{last}")
status_rules(ws, f"H2:H{last}")
ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator="equal", formula=['"High"'], fill=PatternFill("solid", bgColor="F6CFCB")))
ws.conditional_formatting.add(f"G2:G{last}", CellIsRule(operator="equal", formula=['"Medium"'], fill=PatternFill("solid", bgColor="FCE9B8")))

# ---- Sheet: Sections to Add ----
wa = wb.create_sheet("Sections to Add")
cols = ["Proposed #", "Section title", "Where it goes", "Why it is needed", "Draft text", "Check or fill in before using", "Source in project files", "Status"]
header(wa, cols, [13, 26, 24, 34, 100, 40, 30, 13])
for a in adds:
    wa.append(list(a) + ["To do"])
    for c in wa[wa.max_row]:
        c.font, c.alignment, c.border = body, wrap, border
    wa.cell(row=wa.max_row, column=2).font = bold
alast = wa.max_row
wa.auto_filter.ref = f"A1:H{alast}"
dv2 = DataValidation(type="list", formula1=STATUSES, allow_blank=True)
wa.add_data_validation(dv2); dv2.add(f"H2:H{alast}")
status_rules(wa, f"H2:H{alast}")

# ---- Sheet: Read Me ----
wr = wb.create_sheet("Read Me", 0)
wr.column_dimensions["A"].width = 30
wr.column_dimensions["B"].width = 14
wr.column_dimensions["C"].width = 14
wr.column_dimensions["D"].width = 14
wr.column_dimensions["E"].width = 14
wr.column_dimensions["F"].width = 14
wr["A1"] = "IQP Report Revision Tracker"
wr["A1"].font = Font(name=FONT, size=16, bold=True)
wr["A2"] = "Interference Patterns: A Resonant Space - draft dated May 4, 2026 (37 pages)"
wr["A2"].font = body
notes = [
 "How to use this workbook",
 "Changes by Section: one row per edit to the existing draft. Filter by Section, Type or Priority. Page numbers are the ones printed in the draft.",
 "Sections to Add: proposed new sections with draft text ready to paste and adapt.",
 "Edit only the Status and Notes columns. Status is a dropdown: To do, In progress, Done, Not needed.",
 "Priority: High = content that is out of date or missing. Medium = structure and citation errors. Low = grammar and wording.",
 "",
 "Before you use the draft text",
 "Draft sections were written from the project files (firmware, notes and datasheets dated 21-23 September 2026). Anything that has changed since then must be updated.",
 "Text in [square brackets] is a placeholder for a fact only you have, such as test results or actuator counts. Do not submit it as written.",
 "The project notes say the revised motor board was not yet built and the gateway was not yet run on hardware. Section 4.3 must report what was actually tested.",
 "The report describes an arch wall with scissor units; the control spec describes 13 actuators in a 5 / 3 / 5 honeycomb. Reconcile these before using 3.4.4.",
 "Grammar rows come from the extracted text of the PDF. Figures and appendix images were not reviewed.",
]
r = 4
for t in notes:
    c = wr.cell(row=r, column=1, value=t)
    c.font = bold if t in ("How to use this workbook", "Before you use the draft text") else body
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if t and c.font is body:
        wr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        wr.row_dimensions[r].height = 30
    r += 1

r += 1
wr.cell(row=r, column=1, value="Progress").font = bold
r += 1
heads = ["Type", "Total", "To do", "In progress", "Done", "Not needed"]
for i, h in enumerate(heads, 1):
    c = wr.cell(row=r, column=i, value=h)
    c.font, c.fill, c.border = hdr_font, hdr_fill, border
first = r + 1
CS = "'Changes by Section'"
for typ in [C, S, R, G]:
    r += 1
    wr.cell(row=r, column=1, value=typ)
    wr.cell(row=r, column=2, value=f"=COUNTIF({CS}!$D$2:$D${last},$A{r})")
    for j, st in enumerate(["To do", "In progress", "Done", "Not needed"], 3):
        wr.cell(row=r, column=j, value=f'=COUNTIFS({CS}!$D$2:$D${last},$A{r},{CS}!$H$2:$H${last},"{st}")')
r += 1
wr.cell(row=r, column=1, value="New sections to add")
wr.cell(row=r, column=2, value=f"=COUNTA('Sections to Add'!$B$2:$B${alast})")
for j, st in enumerate(["To do", "In progress", "Done", "Not needed"], 3):
    wr.cell(row=r, column=j, value=f"=COUNTIF('Sections to Add'!$H$2:$H${alast},\"{st}\")")
r += 1
wr.cell(row=r, column=1, value="Total")
for j in range(2, 7):
    L = get_column_letter(j)
    wr.cell(row=r, column=j, value=f"=SUM({L}{first}:{L}{r-1})")
for row in wr.iter_rows(min_row=first, max_row=r, max_col=6):
    for c in row:
        c.font = bold if c.row == r else body
        c.border = border

wb.save("IQP_Report_Revision_Tracker.xlsx")

# ---- Markdown copy of the new-section drafts for the project ----
with open("iqp-report-revision-plan.md", "w") as f:
    f.write("# IQP Report Revision Plan\n\n")
    f.write("> Review of the report draft dated May 4, 2026 (*Interference Patterns: A Resonant Space*), made 2026-10-05.\n")
    f.write("> Full row-by-row tracker (grammar, citations, structure): `IQP_Report_Revision_Tracker.xlsx`, delivered in chat.\n")
    f.write("> Draft text below was written from the project files as of 21-23 Sept 2026. Text in [brackets] is a placeholder.\n\n")
    f.write("## Content and structure changes to the existing draft\n\n| Section | Page | Type | Issue | Change |\n|---|---|---|---|---|\n")
    for sec, typ, cur, fix in changes:
        if typ != G:
            f.write(f"| {sec} | {PAGES[sec]} | {typ} | {cur.replace('|','/')} | {fix.replace('|','/')} |\n")
    f.write(f"\n({sum(1 for c in changes if c[1]==G)} grammar and wording fixes are in the spreadsheet only.)\n\n## Sections to add\n\n")
    for num, title, where, why, draft, check, src in adds:
        f.write(f"### {num} {title}\n\n- **Where:** {where}\n- **Why:** {why}\n- **Check before using:** {check}\n- **Source:** {src}\n\n{draft}\n\n")
print(len(changes), len(adds), last, alast)
